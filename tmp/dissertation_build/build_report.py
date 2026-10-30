from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('output/dissertation').resolve()
SOURCE=ROOT/'final_dissertation.md'
OUT=ROOT/'Bridging_the_Chasm_Final_Report.docx'
MAPFILE=Path('tmp/dissertation_build/page_map.json')
PAGE_MAP=json.loads(MAPFILE.read_text()) if MAPFILE.exists() else {}
TITLE='Bridging the Chasm: A Tri-Phase Framework for Enterprise AI (Corporate Data Transformation, Tuning for Compliance, and Robustness Auditing)'
AUTHOR='Mohammed Alaa Marzouq Sallout'
STUDENT='2024tm93522'
SUPERVISOR='Saad Rauf Ghauri'
EXAMINER='Hasan Buttar'
ABSTRACT='''This dissertation develops an application framework for preparing corporate policy and repository data and adapting Qwen3-14B for coding and a limited set of enterprise questions. The implementation combines traceable conversational records with a pinned native NVIDIA selection, validates source and split boundaries, and applies limited gated-MLP channel pruning followed by BF16 LoRA recovery and NF4 export. The prepared enterprise dataset contains 206 records. A completed ten-step RunPod experiment reduced the reported parameter count from 14,768,307,200 to 14,296,448,000 and produced a reloadable 9,696,947,260-byte export. A coding-focused selection of 10,000 OpenCodeInstruct records passed approved-selection validation; its larger recovery run is still represented by progress logs. Independent coding, enterprise-policy, efficiency and robustness comparisons are pending. The demonstrated result is a functioning data-to-compressed-model workflow; claims of improved coding or reliable policy behaviour require those measurements.'''


def clean(s):return s.replace('**','').replace('`','')


def is_caption(line):
    return bool(re.match(r'^(Table|Figure) [A-Z0-9]+\.\d+ ',line)) and not line.endswith(('.',':','?')) and len(line)<150


def bookmark_id(s):return 'b'+hashlib.sha1(s.encode()).hexdigest()[:20]


counter=0

def anchor(p,text):
    global counter
    counter+=1;a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(counter));a.set(qn('w:name'),bookmark_id(text))
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(counter))
    p._p.insert(0,a);p._p.append(b)


def run_font(run,size=None,bold=None):
    run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for name in ('ascii','hAnsi','cs','eastAsia'):fonts.set(qn('w:'+name),'Times New Roman')
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold


def paragraph(text='',style=None,size=None,align=None,bold=False,spacing=None,after=None):
    p=d.add_paragraph(style=style)
    r=p.add_run(clean(text));run_font(r,size,bold)
    if align is not None:p.alignment=align
    if spacing is not None:p.paragraph_format.line_spacing=spacing
    if after is not None:p.paragraph_format.space_after=Pt(after)
    return p


def blank(height=12):
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(height);p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(0);p.add_run().font.size=Pt(1)


def page_number(section,fmt,start):
    section.footer.is_linked_to_previous=False
    p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);p._p.append(field)
    pg=section._sectPr.find(qn('w:pgNumType'))
    if pg is None:pg=OxmlElement('w:pgNumType');section._sectPr.append(pg)
    pg.set(qn('w:fmt'),fmt);pg.set(qn('w:start'),str(start))


def configure(section):
    section.page_width=Inches(9);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    section.header_distance=section.footer_distance=Inches(.45)


def heading(text,level=1):
    p=paragraph(text,style='Heading '+str(level))
    for r in p.runs:run_font(r,14 if level==1 else 12,True)
    p.paragraph_format.keep_with_next=True
    anchor(p,text);return p


def front_title(text):
    p=paragraph(text,size=14,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.2,after=14)
    anchor(p,text);return p


def table(headers,rows,widths=None,font_size=10.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [7/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=clean(value)
    for row in rows:
        cells=t.add_row().cells
        for c,value in zip(cells,row):c.text=clean(str(value))
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for n,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if n==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side,val in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),val);node.set(qn('w:type'),'dxa');margins.append(node)
            pr.append(margins)
            if n==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.keep_with_next=n==0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(headers)<=2 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:run_font(r,font_size,n==0)
    paragraph('',spacing=1,after=3).paragraph_format.keep_with_next=True
    return t


def contents_entry(text,level=1):
    p=d.add_paragraph()
    p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.left_indent=Inches(.2 if level==2 else 0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.95-(.2 if level==2 else 0)),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    # Real internal link; the displayed page number is refreshed from the rendered PDF.
    h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),bookmark_id(text))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Times New Roman');f.set(qn('w:hAnsi'),'Times New Roman');pr.append(f)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'22');pr.append(size);r.append(pr)
    tx=OxmlElement('w:t');tx.text=text;r.append(tx);h.append(r);p._p.append(h)
    rr=p.add_run('\t'+str(PAGE_MAP.get(text,'00')));run_font(rr,11)
    return p


def cover(inner=False):
    blank(14)
    paragraph('A REPORT',size=14,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1,after=14)
    paragraph('ON',size=12,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1,after=14)
    paragraph(TITLE.upper(),style='Title',size=14,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.4,after=18)
    paragraph('BY',size=12,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1,after=12)
    paragraph(AUTHOR,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.2,after=3)
    paragraph('ID No. '+STUDENT,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.2,after=3)
    if inner:
        paragraph('M. Tech Software Engineering - Full Stack Engineering',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.2,after=16)
        paragraph('Prepared in partial fulfilment of the WILP Dissertation Course\nSE-ZG628T',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.4,after=12)
    else:blank(13)
    paragraph('AT',size=12,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1,after=12)
    paragraph('Seer Solutions DMCC\nDubai, United Arab Emirates',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.4,after=18)
    if inner:
        paragraph('Under the supervision of\n'+SUPERVISOR,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.3,after=16)
    blank(10)
    paragraph('BIRLA INSTITUTE OF TECHNOLOGY & SCIENCE, PILANI',size=12,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.2,after=8)
    paragraph('Month and year: ____________________',size=11,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1,after=0)


source=SOURCE.read_text()
heads=[];captions=[]
for line in source.splitlines():
    if line.startswith('# '):heads.append((line[2:],1))
    elif line.startswith('## '):heads.append((line[3:],2))
    elif is_caption(line):captions.append(line)


d=Document();configure(d.sections[0])
for style in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
    s=d.styles[style];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(12);s.paragraph_format.widow_control=True
normal=d.styles['Normal'];normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ('Heading 1','Heading 2','Heading 3'):
    s=d.styles[name];s.paragraph_format.line_spacing=1.2;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(8)
    s.paragraph_format.keep_with_next=True
for r in d.sections[0].footer.paragraphs[0].runs:r.text=''
d.core_properties.title=TITLE;d.core_properties.author=AUTHOR;d.core_properties.subject='Dissertation on corporate data transformation and Qwen3 model compression'
d.core_properties.keywords='Qwen3; corporate policy; coding; LoRA; structured pruning; quantization'
cover(False)
front=d.add_section(WD_SECTION_START.NEW_PAGE);configure(front);page_number(front,'lowerRoman',1)
cover(True)
d.add_page_break();front_title('Acknowledgements')
paragraph('I thank Saad Rauf Ghauri for supervising this dissertation at Seer Solutions DMCC. I also acknowledge Hasan Buttar in his role as Additional Examiner.')
paragraph('The implementation uses open-source software and published research from the Qwen team, NVIDIA, Hugging Face and other contributors. The references identify the methods and datasets used. Responsibility for the implementation, interpretation of its evidence and remaining experimental work rests with the author.')
paragraph('Faculty Mentor: ______________________________',spacing=1.5,after=18)
paragraph(AUTHOR+'\n'+STUDENT,align=WD_ALIGN_PARAGRAPH.RIGHT,spacing=1.4)
d.add_page_break();front_title('Abstract sheet')
for text in [
 'BIRLA INSTITUTE OF TECHNOLOGY & SCIENCE, PILANI (RAJASTHAN) | WILP Division',
 'Organisation: Seer Solutions DMCC    Location: Dubai, UAE',
 'Duration: ____________________    Date of start: ____________________',
 'Date of submission: ____________________',
 'Project: '+TITLE,
 'Student: '+AUTHOR+'    ID: '+STUDENT,
 'Supervisor: '+SUPERVISOR+', Director - Solution Engineering Development',
 'Additional Examiner: '+EXAMINER+', Principal Solution Architect Development',
 'Faculty Mentor: ______________________________',
 'Project areas: Software engineering; machine learning; LLM operations',
 'Keywords: Enterprise data; coding; Qwen3; pruning; LoRA; quantization']:
 paragraph(text,size=10.5,spacing=1.05,after=3,align=WD_ALIGN_PARAGRAPH.LEFT)
paragraph('Abstract',size=12,bold=True,spacing=1,after=5)
paragraph(ABSTRACT,size=12,spacing=2,after=8)
paragraph('Student signature: __________________    Supervisor signature: __________________',size=10.5,spacing=1.1,after=3)
paragraph('Date: __________________                  Date: __________________',size=10.5,spacing=1.1,after=0)
d.add_page_break();front_title('Supervisor certificate')
paragraph('Dissertation title: '+TITLE)
paragraph('Student: '+AUTHOR+' ('+STUDENT+')\nCourse: SE-ZG628T, Dissertation')
paragraph('The dissertation has been prepared at Seer Solutions DMCC under the supervision of '+SUPERVISOR+'. The completed report is to be reviewed for the WILP dissertation requirements before the supervisor records approval for submission.')
paragraph('Completion note F1: supervisor endorsement and signature remain pending. Final certification should follow completion and review of the outstanding evaluation work.',size=11)
blank(28)
for text in ['Supervisor: '+SUPERVISOR,'Designation: Director - Solution Engineering Development','Organisation: Seer Solutions DMCC, Dubai, UAE','Signature: ______________________________','Date: ____________________    Place: ____________________']:
 paragraph(text,spacing=1.5,after=8)
d.add_page_break();front_title('Contents')
for title in ('Acknowledgements','Abstract sheet','Supervisor certificate','List of figures and tables','Abbreviations'):
 contents_entry(title)
for text,level in heads:contents_entry(text,level)
contents_entry('Final report checklist')
d.add_page_break();front_title('List of figures and tables')
for prefix in ('Figure','Table'):
 paragraph(prefix+'s',size=12,bold=True,spacing=1.1,after=7)
 for text in captions:
  if text.startswith(prefix+' '):contents_entry(text)
d.add_page_break();front_title('Abbreviations')
table(['Abbreviation','Expansion'],[
 ('AST','Abstract syntax tree'),('BF16','Brain floating point, 16-bit representation'),('CLI','Command-line interface'),
 ('CUDA','Compute Unified Device Architecture'),('GPU','Graphics processing unit'),('HR','Human resources'),
 ('JSONL','JavaScript Object Notation Lines'),('LLM','Large language model'),('LoRA','Low-Rank Adaptation'),
 ('MLP','Multilayer perceptron'),('NF4','NormalFloat 4-bit'),('PEFT','Parameter-Efficient Fine-Tuning'),
 ('QLoRA','Quantized Low-Rank Adaptation'),('RNG','Random number generator'),('SFT','Supervised fine-tuning'),
 ('VRAM','Video random-access memory')],widths=[1.3,5.7],font_size=11)

body=d.add_section(WD_SECTION_START.NEW_PAGE);configure(body);page_number(body,'decimal',1)
lines=source.splitlines();i=0;first_heading=True;pending_image_caption=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        new_page=not first_heading
        first_heading=False;p=heading(line[2:],1);p.paragraph_format.page_break_before=new_page;i+=1;continue
    if line.startswith('## '):heading(line[3:],2);i+=1;continue
    if line.startswith('```'):
        code=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        for number,text in enumerate(code):
            p=paragraph(text or ' ',size=10,spacing=1.05,after=1,align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next=number<len(code)-1
            p.paragraph_format.keep_together=True
            for r in p.runs:r.font.name='Courier New';r._element.rPr.rFonts.set(qn('w:ascii'),'Courier New');r._element.rPr.rFonts.set(qn('w:hAnsi'),'Courier New')
        i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in row):block.append(row)
            i+=1
        widths=None
        if len(block[0])==2:widths=[2.8,4.2]
        elif len(block[0])==3:widths=[1.6,2.7,2.7]
        elif len(block[0])==4:widths=[2.5,1.5,1.5,1.5]
        if block[0]==['Source','Unique examples','Training rows','Conversation tokens']:widths=[2.3,1.4,1.4,1.9]
        if block[0]==['Condition','Coding execution score','Policy score','Latency and VRAM']:widths=[2.3,1.7,1.2,1.8]
        if block[0]==['Note','Required completion','Evidence needed']:widths=[.6,2.8,3.6]
        if block[0]==['ID','Evidence','Interpretation']:widths=[.6,2.8,3.6]
        if block[0]==['Term','Meaning in this report']:
            block[0].append('First page')
            for row in block[1:]:row.append(str(PAGE_MAP.get('TERM:'+row[0],'00')))
            widths=[1.35,4.95,.7]
        table(block[0],block[1:],widths=widths);continue
    m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if m:
        p=d.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1
        p.add_run().add_picture(str(ROOT/m.group(2)),width=Inches(7))
        for shape in p._p.xpath('.//wp:docPr'):shape.set('descr',m.group(1))
        i+=1;continue
    if is_caption(line):
        p=paragraph(line,size=11,spacing=1.15,after=6)
        p.paragraph_format.keep_with_next=line.startswith('Table')
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if line.startswith('Table') else WD_ALIGN_PARAGRAPH.CENTER
        anchor(p,line);i+=1;continue
    text=line
    while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].startswith(('#','|','```','![')):
        text+=' '+lines[i+1].strip();i+=1
    if text.startswith('Completion note '):
        p=paragraph(text,size=11)
        for r in p.runs:r.italic=True
        d.add_comment(p.runs,text=text,author='Completion notes',initials='CN')
    elif text.startswith('[') and re.match(r'^\[\d+\]',text):
        p=paragraph(text,size=11,spacing=1.25,after=10,align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.keep_together=True
    else:paragraph(text)
    i+=1

# The student's verification sheet is the LAST page; signatures are never fabricated.
d.add_page_break();front_title('Final report checklist')
paragraph('Checklist of items for the Final Dissertation / Project / Project Work Report',size=10.5,bold=True,spacing=1,after=5)
paragraph('Complete, verify and sign this sheet after the outstanding work and supervisor review.',size=10.5,spacing=1,after=6)
items=[
 ('1','Technical-report elements and neat formatting','Final review'),
 ('2','Cover follows the supplied Appendix A specimen','Prepared'),
 ('3','Inner title page follows Appendix B','Prepared'),
 ('4(a)','Supervisor certificate included for approval','Review pending'),
 ('4(b)','Supervisor certificate signed','Pending'),
 ('5','One-page abstract, at most 200 words, with keywords','Prepared'),
 ('6','Title reflects the completed scope','Final scope review'),
 ('7','Abbreviations and acronyms included','Prepared'),
 ('8','Literature survey included','Prepared'),
 ('9','Contents, page numbers, figures, tables, captions and appendices checked','Final verification'),
 ('10','Conclusions follow the completed results','Pending evaluation'),
 ('11','References included and cited in the text','Prepared'),
 ('12','Format and content meet the guidelines; no source-code dump','Final review')]
table(['Item','Verification','Status'],items,widths=[.55,4.9,1.55],font_size=10)
paragraph('Declaration by student',size=10.5,bold=True,spacing=1,after=4)
paragraph('I certify that I have verified all items in this checklist and that the completed report meets the format specified in the course handout.',size=10.5,spacing=1.1,after=9)
paragraph('Place: ____________________    Student signature: ____________________',size=10.5,spacing=1.15,after=8)
paragraph('Date: ____________________    Name: '+AUTHOR,size=10.5,spacing=1.15,after=8)
paragraph('ID No.: '+STUDENT,size=10.5,spacing=1.15,after=0)

# Avoid automatic field-refresh prompts that can replace the checked static contents.
settings=d.settings._element
update=settings.find(qn('w:updateFields'))
if update is not None:settings.remove(update)
d.save(OUT)
(ROOT/'document_metadata.json').write_text(json.dumps({'title':TITLE,'author':AUTHOR,'student_id':STUDENT,
 'abstract_word_count':len(ABSTRACT.split()),'body_source_word_count':len(source.split()),
 'paper_inches':[9,11],'margins_inches':1,'body_font':'Times New Roman 12pt','body_line_spacing':2,
 'pending_faculty_mentor':True,'pending_submission_date':True,'pending_signatures':True,
 'headings':heads,'captions':captions,'output_docx':str(OUT)},indent=2))
print(OUT)
